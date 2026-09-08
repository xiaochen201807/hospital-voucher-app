#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
根据西药药品入库单与总账/辅助核算表自动生成入库凭证导入模板
规范参考：凭证导入模板-入库.xlsx

生成逻辑：
1. 读取【西药-药品入库单.xlsx】中的入库明细（药品名称、规格、数量、进价、进价金额、供应商等）。
2. 从凭证导入模板 Sheet 3【辅助核算项目数据】读取供应商字典与存货字典，从总账表读取存货科目与单价。
3. 按照【供应商】对入库明细进行分组，为每个供应商生成：
   - 借方存货明细行（科目代码 1201）：
     * 摘要：{供应商简称}到货（例如：国药乐仁堂到货）
     * 科目代码：1201
     * 借方金额：药品进价金额
     * 贷方金额：留空
     * 供应商：留空
     * 存货：匹配到的存货辅助核算编码（例如 XY0015）
     * 数量：入库数量
     * 单价：入库进价
     * 原币金额：进价金额
   - 贷方应付账款汇总行（科目代码 220201）：
     * 摘要：{供应商简称}到货（例如：国药乐仁堂到货）
     * 科目代码：220201
     * 借方金额：留空
     * 贷方金额：该供应商名下所有药品进价金额之和
     * 供应商：该供应商对应的辅助核算编码（例如 001）
     * 存货：留空
     * 原币金额：贷方汇总金额
   - 全局分录序号从 1 递增，日期固定为月末最后一天（或指定日期）。
"""

import os
import sys
import re
import json
import calendar
import argparse
from datetime import datetime, date
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side


def load_factory_config(config_file='factory_mapping.json'):
    """
    读取厂家全称与简写映射配置文件
    """
    if not os.path.exists(config_file):
        return {
            'factory_abbreviations': {},
            'strict_vendor_suffix_drugs': [],
            'drug_code_overrides': {}
        }
    try:
        with open(config_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return {
                'factory_abbreviations': data.get('factory_abbreviations', {}),
                'strict_vendor_suffix_drugs': data.get('strict_vendor_suffix_drugs', []),
                'drug_code_overrides': data.get('drug_code_overrides', {})
            }
    except Exception as e:
        print(f"[!] 读取厂家映射文件 {config_file} 失败: {e}")
        return {
            'factory_abbreviations': {},
            'strict_vendor_suffix_drugs': [],
            'drug_code_overrides': {}
        }


def normalize_text(text):
    """文本标准化：去除所有空格，统一半角全角括号与符号"""
    if not text:
        return ''
    s = str(text).strip()
    s = re.sub(r'\s+', '', s)
    s = s.replace('（', '(').replace('）', ')')
    s = s.replace('【', '[').replace('】', ']')
    s = s.replace('×', '*').replace('x', '*').replace('X', '*')
    return s


def detect_month_end_date(inbound_file=None, ledger_file=None, max_inbound_date=None):
    """
    自动检测目标年月并返回月末最后一天的日期字符串（YYYY-MM-DD）
    """
    # 优先从入库明细中的最大日期推导
    if max_inbound_date:
        m = re.search(r'(\d{4})[/-](\d{1,2})[/-](\d{1,2})', str(max_inbound_date))
        if m:
            year, month = int(m.group(1)), int(m.group(2))
            last_day = calendar.monthrange(year, month)[1]
            return f"{year:04d}-{month:02d}-{last_day:02d}"

    # 其次从文件名提取，例如 "2026.8月", "202608期", "2026-08"
    combined_names = f"{os.path.basename(inbound_file or '')} {os.path.basename(ledger_file or '')}"
    
    m = re.search(r'(\d{4})[^\d]*?(\d{1,2})月', combined_names)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
        last_day = calendar.monthrange(year, month)[1]
        return f"{year:04d}-{month:02d}-{last_day:02d}"

    m = re.search(r'(\d{4})(\d{2})期?', combined_names)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
        if 1 <= month <= 12:
            last_day = calendar.monthrange(year, month)[1]
            return f"{year:04d}-{month:02d}-{last_day:02d}"

    # 默认当前时间所在月末
    today = date.today()
    last_day = calendar.monthrange(today.year, today.month)[1]
    return f"{today.year:04d}-{today.month:02d}-{last_day:02d}"


def get_supplier_brief(sup_name):
    """
    提取供应商简称并拼接‘到货’摘要
    例如：'国药乐仁堂医药有限公司' -> '国药乐仁堂到货'
    """
    if not sup_name:
        return '药品到货'
    clean = re.sub(r'(股份有限公司|医药有限公司|药业有限公司|有限公司|有限责任公司|责任公司)', '', str(sup_name).strip())
    return f"{clean}到货"


def load_inbound_data(inbound_file):
    """
    读取药品入库单文件（.xlsx 或 .xls）。
    自动识别表头行，提取明细记录。
    返回: (sheet_name, inbound_items, max_date_str)
    """
    ext = os.path.splitext(inbound_file)[1].lower()
    inbound_items = []
    max_date_str = None

    if ext == '.xlsx':
        wb = openpyxl.load_workbook(inbound_file, data_only=True)
        target_sheet = None
        for s in wb.worksheets:
            for r in range(1, min(10, s.max_row + 1)):
                row_vals = [str(s.cell(r, c).value or '').strip() for c in range(1, s.max_column + 1)]
                if '药品名称' in row_vals and ('数量' in row_vals or '进价' in row_vals):
                    target_sheet = s
                    header_row = r
                    break
            if target_sheet:
                break

        if not target_sheet:
            target_sheet = wb.worksheets[0]
            header_row = 4

        sheet_name = target_sheet.title
        header_vals = [str(target_sheet.cell(header_row, c).value or '').strip() for c in range(1, target_sheet.max_column + 1)]
        
        def find_col(col_name):
            for idx, val in enumerate(header_vals, start=1):
                if col_name in val:
                    return idx
            return -1

        col_order_no = find_col('入库单号')
        col_code = find_col('药品编码')
        col_date = find_col('入库日期')
        col_name = find_col('药品名称')
        col_factory = find_col('厂家')
        if col_factory < 0:
            col_factory = find_col('制药厂')
        if col_factory < 0:
            col_factory = find_col('生产企业')
        col_unit = find_col('单位')
        col_spec = find_col('规格')
        col_qty = find_col('数量')
        col_batch = find_col('批号')
        col_price = find_col('进价')
        col_amt = find_col('进价金额')
        col_supplier = find_col('供应商')

        for r in range(header_row + 1, target_sheet.max_row + 1):
            raw_name = str(target_sheet.cell(r, col_name).value or '').strip() if col_name > 0 else ''
            if not raw_name or '合计' in raw_name or '制表' in raw_name:
                continue

            spec = str(target_sheet.cell(r, col_spec).value or '').strip() if col_spec > 0 else ''
            unit = str(target_sheet.cell(r, col_unit).value or '').strip() if col_unit > 0 else ''
            order_no = str(target_sheet.cell(r, col_order_no).value or '').strip() if col_order_no > 0 else ''
            drug_code = str(target_sheet.cell(r, col_code).value or '').strip() if col_code > 0 else ''
            date_val = str(target_sheet.cell(r, col_date).value or '').strip() if col_date > 0 else ''
            batch_no = str(target_sheet.cell(r, col_batch).value or '').strip() if col_batch > 0 else ''
            supplier = str(target_sheet.cell(r, col_supplier).value or '').strip() if col_supplier > 0 else ''
            factory = str(target_sheet.cell(r, col_factory).value or '').strip() if col_factory > 0 else ''

            if date_val and (max_date_str is None or date_val > max_date_str):
                max_date_str = date_val

            try:
                raw_qty = float(target_sheet.cell(r, col_qty).value or 0.0) if col_qty > 0 else 0.0
                qty = int(raw_qty) if raw_qty.is_integer() else raw_qty
            except (ValueError, TypeError):
                qty = 0.0

            try:
                price = float(target_sheet.cell(r, col_price).value or 0.0) if col_price > 0 else 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                amt = float(target_sheet.cell(r, col_amt).value or 0.0) if col_amt > 0 else 0.0
            except (ValueError, TypeError):
                amt = round(qty * price, 2)

            inbound_items.append({
                'row_index': r,
                'order_no': order_no,
                'drug_code': drug_code,
                'date': date_val,
                'name': raw_name,
                'spec': spec,
                'unit': unit,
                'qty': qty,
                'in_price': price,
                'in_amt': amt,
                'batch_no': batch_no,
                'supplier': supplier,
                'factory': factory or supplier
            })

    elif ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(inbound_file)
        target_sheet = wb.sheet_by_index(0)
        sheet_name = target_sheet.name

        header_row = None
        for r in range(min(10, target_sheet.nrows)):
            row_vals = [str(target_sheet.cell_value(r, c)).strip() for c in range(target_sheet.ncols)]
            if '药品名称' in row_vals and ('数量' in row_vals or '进价' in row_vals):
                header_row = r
                break

        if header_row is None:
            header_row = 3

        header_vals = [str(target_sheet.cell_value(header_row, c)).strip() for c in range(target_sheet.ncols)]

        def find_xls_col(col_name):
            for idx, val in enumerate(header_vals):
                if col_name in val:
                    return idx
            return -1

        col_order_no = find_xls_col('入库单号')
        col_code = find_xls_col('药品编码')
        col_date = find_xls_col('入库日期')
        col_name = find_xls_col('药品名称')
        col_factory = find_xls_col('厂家')
        if col_factory < 0:
            col_factory = find_xls_col('制药厂')
        if col_factory < 0:
            col_factory = find_xls_col('生产企业')
        col_spec = find_xls_col('规格')
        col_unit = find_xls_col('单位')
        col_qty = find_xls_col('数量')
        col_batch = find_xls_col('批号')
        col_price = find_xls_col('进价')
        col_amt = find_xls_col('进价金额')
        col_supplier = find_xls_col('供应商')

        for r in range(header_row + 1, target_sheet.nrows):
            raw_name = str(target_sheet.cell_value(r, col_name)).strip() if col_name >= 0 else ''
            if not raw_name or '合计' in raw_name or '制表' in raw_name:
                continue

            spec = str(target_sheet.cell_value(r, col_spec)).strip() if col_spec >= 0 else ''
            unit = str(target_sheet.cell_value(r, col_unit)).strip() if col_unit >= 0 else ''
            order_no = str(target_sheet.cell_value(r, col_order_no)).strip() if col_order_no >= 0 else ''
            drug_code = str(target_sheet.cell_value(r, col_code)).strip() if col_code >= 0 else ''
            date_val = str(target_sheet.cell_value(r, col_date)).strip() if col_date >= 0 else ''
            batch_no = str(target_sheet.cell_value(r, col_batch)).strip() if col_batch >= 0 else ''
            supplier = str(target_sheet.cell_value(r, col_supplier)).strip() if col_supplier >= 0 else ''
            factory = str(target_sheet.cell_value(r, col_factory)).strip() if col_factory >= 0 else ''

            if date_val and (max_date_str is None or date_val > max_date_str):
                max_date_str = date_val

            try:
                raw_qty = float(target_sheet.cell_value(r, col_qty) or 0.0) if col_qty >= 0 else 0.0
                qty = int(raw_qty) if raw_qty.is_integer() else raw_qty
            except (ValueError, TypeError):
                qty = 0.0

            try:
                price = float(target_sheet.cell_value(r, col_price) or 0.0) if col_price >= 0 else 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                amt = float(target_sheet.cell_value(r, col_amt) or 0.0) if col_amt >= 0 else 0.0
            except (ValueError, TypeError):
                amt = round(qty * price, 2)

            inbound_items.append({
                'row_index': r + 1,
                'order_no': order_no,
                'drug_code': drug_code,
                'date': date_val,
                'name': raw_name,
                'spec': spec,
                'unit': unit,
                'qty': qty,
                'in_price': price,
                'in_amt': amt,
                'batch_no': batch_no,
                'supplier': supplier,
                'factory': factory or supplier
            })
    else:
        raise ValueError(f"不支持的文件格式: {ext}，仅支持 .xlsx 或 .xls")

    return sheet_name, inbound_items, max_date_str


def load_ledger_and_aux_data(ledger_file, template_file):
    """
    读取总账表及凭证模板中的辅助核算项目数据（存货字典与供应商字典）
    返回: (ledger_exact, ledger_norm, aux_codes_set, aux_name_map, supplier_dict)
    """
    ledger_exact = {}
    ledger_norm = {}

    if ledger_file and os.path.exists(ledger_file):
        wb_ledger = openpyxl.load_workbook(ledger_file, data_only=True)
        s_ledger = wb_ledger.active

        for r in range(4, s_ledger.max_row + 1):
            c1 = str(s_ledger.cell(r, 1).value or '').strip()
            c2 = str(s_ledger.cell(r, 2).value or '').strip()
            if not c2:
                continue

            aux_code = c1.split('_')[-1] if '_' in c1 else c1

            # 提取单价 (优先取第 18 列【期末结存单价】，其次取第 6 列【期初单价】)
            def parse_price(val):
                try:
                    p = float(val) if val is not None and str(val).strip() != '' else 0.0
                    return p if p > 0.0 else None
                except (ValueError, TypeError):
                    return None

            price = parse_price(s_ledger.cell(r, 18).value)
            if price is None:
                price = parse_price(s_ledger.cell(r, 6).value)

            item_info = {
                'subject_code': c1,
                'aux_code': aux_code,
                'subject_name': c2,
                'price': price
            }

            ledger_exact[c2] = item_info
            ledger_norm[normalize_text(c2)] = item_info

    # 读取凭证导入模板中的【辅助核算项目数据】(Sheet 3)
    aux_codes_set = set()
    aux_name_map = {}
    supplier_dict = {} # code -> name
    if template_file and os.path.exists(template_file):
        wb_tmpl = openpyxl.load_workbook(template_file, data_only=True)
        if '辅助核算项目数据' in wb_tmpl.sheetnames:
            s_aux = wb_tmpl['辅助核算项目数据']
            for r in range(3, s_aux.max_row + 1):
                cat = str(s_aux.cell(r, 1).value or '').strip()
                code = str(s_aux.cell(r, 2).value or '').strip()
                name = str(s_aux.cell(r, 3).value or '').strip()
                if cat == '存货':
                    if code:
                        aux_codes_set.add(code)
                    if name:
                        aux_name_map[normalize_text(name)] = code
                elif cat == '供应商':
                    if code and name:
                        supplier_dict[code] = name

    return ledger_exact, ledger_norm, aux_codes_set, aux_name_map, supplier_dict


def match_supplier_code(sup_name, supplier_dict):
    """
    将入库单中的供应商名称匹配到辅助核算项目数据中的供应商编码（例如 '001'）
    """
    if not sup_name or not supplier_dict:
        return '', ''

    norm_in = normalize_text(sup_name)

    # 1. 精确匹配
    for c, n in supplier_dict.items():
        if normalize_text(n) == norm_in:
            return c, n

    # 2. 互相包含匹配
    for c, n in supplier_dict.items():
        norm_n = normalize_text(n)
        if norm_n in norm_in or norm_in in norm_n:
            return c, n

    # 3. 去除常见公司后缀后比对核心关键词
    clean_in = re.sub(r'(股份有限公司|医药有限公司|药业有限公司|有限公司|责任公司)', '', norm_in)
    for c, n in supplier_dict.items():
        clean_n = re.sub(r'(股份有限公司|医药有限公司|药业有限公司|有限公司|责任公司)', '', normalize_text(n))
        if clean_in and clean_n and (clean_in in clean_n or clean_n in clean_in):
            return c, n

    return '', ''


def match_drug_item(item, ledger_exact, ledger_norm, aux_name_map, factory_config=None):
    """
    为药品匹配总账科目与辅助核算编码
    规则：优先按 厂家/供应商 简称进行带括号标识匹配（如 鸡血藤（蕴德）、海螵蛸（蕴德））
    若属于严格厂家限定品种（如 海螵蛸），不回退到无厂家的老科目；
    非严格品种且无其他厂家冲突时，回退到标准科目匹配。
    """
    name = item['name']
    spec = item['spec']
    factory = item.get('factory') or item.get('supplier') or ''

    if factory_config is None:
        factory_config = {}
    factory_abbrs = factory_config.get('factory_abbreviations', {})
    strict_drugs = factory_config.get('strict_vendor_suffix_drugs', [])
    overrides = factory_config.get('drug_code_overrides', {})

    # 1. 查找厂家对应的简写（例如 河北蕴德药业有限公 -> 蕴德）
    abbr = None
    if factory:
        norm_fac = normalize_text(factory)
        for full_f, short_f in factory_abbrs.items():
            norm_k = normalize_text(full_f)
            if norm_k == norm_fac or norm_k in norm_fac or norm_fac in norm_k:
                abbr = short_f
                break

    all_vendor_tags = set(factory_abbrs.values()) | {'国瑞堂', '国松堂', '盛方', '神农', '蕴德'}
    other_vendor_tags = {v for v in all_vendor_tags if v != abbr}

    # 2. 如果存在厂家简写，构造带厂家后缀的目标药名优先匹配
    if abbr:
        target_name_with_vendor = f"{name}（{abbr}）"
        target_name_half = f"{name}({abbr})"

        # 2.1 检查是否存在直接指定的编码覆写 (drug_code_overrides)
        for cand_name in [target_name_with_vendor, target_name_half, name]:
            if cand_name in overrides and overrides[cand_name]:
                return {
                    'subject_code': '',
                    'aux_code': overrides[cand_name],
                    'subject_name': target_name_with_vendor,
                    'price': None
                }

        # 2.2 在总账中寻找带厂家简写的科目 (精确与标准化)
        for cand_vendor in [target_name_with_vendor, target_name_half]:
            c1 = f"存货_{cand_vendor} {spec}".strip()
            if c1 in ledger_exact:
                return ledger_exact[c1]
            c2 = f"存货_{cand_vendor}".strip()
            if c2 in ledger_exact:
                return ledger_exact[c2]
            
            n1 = normalize_text(f"存货_{cand_vendor}{spec}")
            if n1 in ledger_norm:
                return ledger_norm[n1]
            n2 = normalize_text(f"存货_{cand_vendor}")
            if n2 in ledger_norm:
                return ledger_norm[n2]

        # 2.3 辅助核算字典中匹配带厂家简称
        for cand_vendor in [target_name_with_vendor, target_name_half]:
            norm_tv = normalize_text(cand_vendor)
            if norm_tv in aux_name_map:
                return {
                    'subject_code': '',
                    'aux_code': aux_name_map[norm_tv],
                    'subject_name': target_name_with_vendor,
                    'price': None
                }

        # 2.4 如果带厂家的没有匹配到：
        # 如果该药属于严格必须带厂家标识的品种（如 海螵蛸），决不回退到老旧不带厂家的科目（防止误匹配到旧批次如 ZY0106）
        if name in strict_drugs:
            item['target_name'] = target_name_with_vendor
            return None

        # 检查总账中是否存在属于【其他已知厂家】的同名科目（例如 鸡血藤（国瑞堂）），若有冲突绝不能误匹配过去
        norm_name = normalize_text(name)
        has_other_vendor_conflict = False
        for k_norm in ledger_norm.keys():
            if norm_name in k_norm:
                if any(normalize_text(other_v) in k_norm for other_v in other_vendor_tags):
                    has_other_vendor_conflict = True
                    break
        if has_other_vendor_conflict:
            item['target_name'] = target_name_with_vendor
            return None

    # 3. 常规通用名匹配（无厂家要求或无冲突的回退匹配）
    cand1 = f"存货_{name} {spec}".strip()
    if cand1 in ledger_exact:
        return ledger_exact[cand1]

    cand2 = f"存货_{name}".strip()
    if cand2 in ledger_exact:
        return ledger_exact[cand2]

    norm1 = normalize_text(f"存货_{name}{spec}")
    if norm1 in ledger_norm:
        return ledger_norm[norm1]

    norm2 = normalize_text(f"存货_{name}")
    if norm2 in ledger_norm:
        return ledger_norm[norm2]

    norm_name = normalize_text(name)
    norm_spec = normalize_text(spec)
    for k_norm, val in ledger_norm.items():
        if norm_name in k_norm:
            # 排除带其他厂家后缀的科目
            if any(normalize_text(other_v) in k_norm for other_v in other_vendor_tags):
                continue
            if not norm_spec or norm_spec in k_norm:
                return val

    if norm_name in aux_name_map:
        return {
            'subject_code': '',
            'aux_code': aux_name_map[norm_name],
            'subject_name': name,
            'price': None
        }

    return None


def generate_inbound_voucher_file(inbound_file, ledger_file, template_file, output_file=None, voucher_date=None, voucher_no=None, config_file='factory_mapping.json'):
    """
    主生成函数：根据入库单生成符合规范的入库凭证导入模板
    按供应商分组，每个供应商包含存货借方明细（1201）与贷方应付账款汇总（220201，带供应商编码）
    """
    # 0. 读取厂家全名与简写映射配置
    factory_config = load_factory_config(config_file)
    if factory_config.get('factory_abbreviations'):
        print(f"[*] 成功加载厂家全名与简写映射表: 共 {len(factory_config['factory_abbreviations'])} 个映射项")

    # 1. 读取入库明细
    sheet_name, inbound_items, max_date_str = load_inbound_data(inbound_file)
    print(f"[*] 成功读取入库明细: 来自【{os.path.basename(inbound_file)}】的【{sheet_name}】")
    print(f"[*] 共读取到 {len(inbound_items)} 条药品入库记录")

    # 2. 确定凭证入账日期
    target_date = voucher_date or detect_month_end_date(inbound_file, ledger_file, max_date_str)

    # 3. 读取总账与辅助核算字典（包含存货与供应商）
    ledger_exact, ledger_norm, aux_codes_set, aux_name_map, supplier_dict = load_ledger_and_aux_data(ledger_file, template_file)
    print(f"[*] 成功加载总账与辅助核算项目字典 (总账科目: {len(ledger_exact)} 项, 辅助核算存货: {len(aux_codes_set)} 项, 供应商: {len(supplier_dict)} 项)")

    # 4. 按供应商对入库明细进行分组
    grouped_items = OrderedDict()
    for item in inbound_items:
        sup = item['supplier'] or '未知供应商'
        if sup not in grouped_items:
            grouped_items[sup] = []
        grouped_items[sup].append(item)

    print(f"[*] 入库明细共涉及 {len(grouped_items)} 个供应商")

    # 5. 加载凭证导入模板
    wb_out = openpyxl.load_workbook(template_file)
    if '凭证模版' not in wb_out.sheetnames:
        raise ValueError(f"凭证导入模板中未找到【凭证模版】工作表！")
    s_voucher = wb_out['凭证模版']

    # 清空原有数据行（保留第 1 行表头）
    if s_voucher.max_row > 1:
        s_voucher.delete_rows(2, s_voucher.max_row - 1)

    # 样式配置
    font_body = Font(name='微软雅黑', size=10, bold=False)
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    thin_border = Side(style='thin', color='D9D9D9')
    border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # 6. 逐个供应商生成分录（借方存货明细 + 贷方应付账款汇总）
    current_row = 2
    entry_seq = 1
    matched_count = 0
    unmatched_items = []
    total_debit_amt = 0.0
    total_credit_amt = 0.0
    total_qty = 0.0

    for sup_name, items in grouped_items.items():
        sup_code, sup_full_name = match_supplier_code(sup_name, supplier_dict)
        summary_text = get_supplier_brief(sup_name)
        sup_debit_total = 0.0

        # 6.1 生成该供应商名下的借方存货明细行（科目代码 1201）
        for item in items:
            matched = match_drug_item(item, ledger_exact, ledger_norm, aux_name_map, factory_config)

            qty = item['qty']
            price = item['in_price']
            amt = round(item['in_amt'], 2)

            total_qty += qty
            sup_debit_total += amt
            total_debit_amt += amt

            if matched and matched.get('aux_code'):
                aux_code = matched['aux_code']
                matched_count += 1
            else:
                aux_code = ''
                unmatched_items.append(item)

            row_values = [
                target_date,      # 1. 日期: 月末最后一天
                '记',             # 2. 凭证字: 固定为“记”
                voucher_no,       # 3. 凭证号: 留空或指定
                None,             # 4. 附件数: 为空
                entry_seq,        # 5. 分录序号: 全局递增
                summary_text,     # 6. 摘要: {供应商简称}到货 (例如 国药乐仁堂到货)
                1201,             # 7. 科目代码: 1201 (库存药品/存货)
                None,             # 8. 科目名称: 为空
                amt,              # 9. 借方金额: 采购进价金额
                None,             # 10. 贷方金额: 为空
                None,             # 11. 客户: 为空
                None,             # 12. 供应商: 借方存货行留空
                None,             # 13. 职员: 为空
                None,             # 14. 项目: 为空
                None,             # 15. 部门: 为空
                aux_code,         # 16. 存货: 辅助核算编码 (例如 XY0015)
                None,             # 17. 是否限定: 为空
                None,             # 18. 自定义类别: 为空
                None,             # 19. 自定义编码: 为空
                None,             # 20. 自定义类别1: 为空
                None,             # 21. 自定义编码1: 为空
                qty,              # 22. 数量: 入库数量
                price,            # 23. 单价: 入库进价
                amt,              # 24. 原币金额: 采购进价金额
                'RMB',            # 25. 币别: 固定为 RMB
                1,                # 26. 汇率: 固定为 1
            ]

            s_voucher.row_dimensions[current_row].height = 19
            for col_idx, val in enumerate(row_values, start=1):
                cell = s_voucher.cell(row=current_row, column=col_idx, value=val)
                cell.font = font_body
                cell.border = border_all
                if col_idx in [9, 10, 24]:
                    cell.alignment = align_right
                    cell.number_format = '#,##0.00'
                elif col_idx in [22, 23]:
                    cell.alignment = align_right
                else:
                    cell.alignment = align_center

            current_row += 1
            entry_seq += 1

        # 6.2 生成该供应商的贷方应付账款汇总行（科目代码 220201）
        sup_credit_amt = round(sup_debit_total, 2)
        total_credit_amt += sup_credit_amt

        credit_row_values = [
            target_date,          # 1. 日期: 月末最后一天
            '记',                 # 2. 凭证字: 固定为“记”
            voucher_no,           # 3. 凭证号: 留空或指定
            None,                 # 4. 附件数: 为空
            entry_seq,            # 5. 分录序号: 全局递增
            summary_text,         # 6. 摘要: {供应商简称}到货 (例如 国药乐仁堂到货)
            220201,               # 7. 科目代码: 220201 (应付账款)
            None,                 # 8. 科目名称: 为空
            None,                 # 9. 借方金额: 为空
            sup_credit_amt,       # 10. 贷方金额: 该供应商借方金额之和
            None,                 # 11. 客户: 为空
            sup_code,             # 12. 供应商: 辅助核算编码 (例如 001)
            None,                 # 13. 职员: 为空
            None,                 # 14. 项目: 为空
            None,                 # 15. 部门: 为空
            None,                 # 16. 存货: 为空
            None,                 # 17. 是否限定: 为空
            None,                 # 18. 自定义类别: 为空
            None,                 # 19. 自定义编码: 为空
            None,                 # 20. 自定义类别1: 为空
            None,                 # 21. 自定义编码1: 为空
            None,                 # 22. 数量: 留空
            None,                 # 23. 单价: 留空
            sup_credit_amt,       # 24. 原币金额: 贷方汇总金额
            'RMB',                # 25. 币别: 固定为 RMB
            1,                    # 26. 汇率: 固定为 1
        ]

        s_voucher.row_dimensions[current_row].height = 19
        for col_idx, val in enumerate(credit_row_values, start=1):
            cell = s_voucher.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_body
            cell.border = border_all
            if col_idx in [9, 10, 24]:
                cell.alignment = align_right
                cell.number_format = '#,##0.00'
            elif col_idx in [22, 23]:
                cell.alignment = align_right
            else:
                cell.alignment = align_center

        current_row += 1
        entry_seq += 1

    # 保存文件
    save_path = output_file or '凭证导入模板_入库_已生成.xlsx'
    wb_out.save(save_path)

    # 7. 打印总结报表
    total_entries = entry_seq - 1
    print("\n" + "="*55)
    print("           入库凭证导入模板生成完成报告")
    print("="*55)
    print(f"凭证入账日期: {target_date}")
    print(f"涉及供应商数: {len(grouped_items)} 个")
    print(f"生成分录总数: {total_entries} 行（借方明细 {len(inbound_items)} 行 + 贷方汇总 {len(grouped_items)} 行）")
    print(f"存货匹配成功: {matched_count} 行 ({matched_count/len(inbound_items)*100:.1f}%)")
    print(f"未匹配/新药行: {len(unmatched_items)} 行 ({len(unmatched_items)/len(inbound_items)*100:.1f}%)")
    print(f"入库数量总和: {total_qty:,.2f}")
    print(f"借方金额总计: {total_debit_amt:,.2f} 元")
    print(f"贷方金额总计: {total_credit_amt:,.2f} 元 (借贷平衡: {'是' if round(total_debit_amt, 2) == round(total_credit_amt, 2) else '否'})")
    print(f"输出保存文件: {save_path}")

    print("\n--- 各供应商分录汇总 ---")
    for sup_name, items in grouped_items.items():
        sup_code, _ = match_supplier_code(sup_name, supplier_dict)
        brief = get_supplier_brief(sup_name)
        subtotal = sum(it['in_amt'] for it in items)
        print(f"  * 供应商: {sup_name} | 编码: {sup_code} | 摘要: {brief} | 药品数: {len(items)} 种 | 贷方金额: {subtotal:,.2f} 元")

    if unmatched_items:
        print("\n[⚠️ 注意] 以下药品在总账/辅助核算中未匹配到存货编码，存货编码暂已安全留空：")
        for u in unmatched_items:
            show_name = u.get('target_name') or u['name']
            print(f"  - 供应商: {u['supplier']} | 药品: {show_name} | 规格: {u['spec']} | 入库数量: {u['qty']} | 进价: {u['in_price']} | 金额: {u['in_amt']} 元")
        print("  提示：上述药品属于新入库且尚未在财务软件建档品种，请在系统中建档后补填辅助核算编码（或在 factory_mapping.json 中指定映射）。")

    return save_path


def main():
    parser = argparse.ArgumentParser(description='根据西药/中药药品入库单与总账/辅助核算表自动生成入库凭证导入模板（按供应商分组与贷方汇总）')
    parser.add_argument('-i', '--inbound', default=None, help='入库单文件路径（默认自动查找*入库*.xlsx/.xls）')
    parser.add_argument('-l', '--ledger', default=None, help='总账表文件路径（默认自动查找数量金额总账*.xlsx）')
    parser.add_argument('-t', '--template', default='凭证导入模板.xlsx', help='凭证导入模板路径（默认凭证导入模板.xlsx）')
    parser.add_argument('-o', '--output', default='凭证导入模板_入库_已生成.xlsx', help='输出凭证文件路径（默认 凭证导入模板_入库_已生成.xlsx）')
    parser.add_argument('-d', '--date', default=None, help='指定凭证日期（格式 YYYY-MM-DD，默认自动推导月末最后一天）')
    parser.add_argument('-v', '--voucher-no', default=None, help='指定凭证号（默认留空）')
    parser.add_argument('-c', '--config', default='factory_mapping.json', help='厂家全称与简写映射配置文件（默认 factory_mapping.json）')

    args = parser.parse_args()

    # 自动搜索入库单
    inbound_file = args.inbound
    if not inbound_file:
        cands = [f for f in os.listdir('.') if f.endswith(('.xlsx', '.xls')) and '入库' in f and not f.startswith('~$') and '已生成' not in f and '模板' not in f]
        if cands:
            inbound_file = cands[0]
        else:
            # 兼容带有模板命名的入库单搜索
            cands = [f for f in os.listdir('.') if f.endswith(('.xlsx', '.xls')) and '入库' in f and not f.startswith('~$') and '已生成' not in f]
            if cands:
                inbound_file = cands[0]
            else:
                print("错误: 未找到入库单文件，请使用 -i 指定入库单路径。")
                sys.exit(1)

    # 自动搜索总账表
    ledger_file = args.ledger
    if not ledger_file:
        cands = [f for f in os.listdir('.') if f.endswith('.xlsx') and '总账' in f and not f.startswith('~$')]
        if cands:
            ledger_file = cands[0]
        else:
            print("错误: 未找到数量金额总账表，请使用 -l 指定总账表路径。")
            sys.exit(1)

    # 检查模板文件
    template_file = args.template
    if not os.path.exists(template_file):
        print(f"错误: 凭证导入模板 '{template_file}' 不存在。")
        sys.exit(1)

    print(f"使用的入库单: {inbound_file}")
    print(f"使用的总账表: {ledger_file}")
    print(f"使用的模板表: {template_file}")
    print(f"使用的映射表: {args.config}")

    generate_inbound_voucher_file(
        inbound_file=inbound_file,
        ledger_file=ledger_file,
        template_file=template_file,
        output_file=args.output,
        voucher_date=args.date,
        voucher_no=args.voucher_no,
        config_file=args.config
    )


if __name__ == '__main__':
    main()
