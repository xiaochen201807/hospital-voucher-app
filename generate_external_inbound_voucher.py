#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
外账入库凭证生成脚本
功能：
1. 读取西药药品入库单（西药入库单.xlsx / 西药-药品入库单.xlsx，支持 .xlsx 和 .xls）；
2. 读取外账导入模板（表格迁账参考模板.xlsx）的【辅助信息】Sheet，提取供应商与存货辅助档案；
3. 按照各供应商对入库药品进行分组：
   - 借方分录（入库明细）：科目 1405 库存商品，填写数量、进价金额、存货辅助编码与名称；
   - 贷方汇总（供应商汇总）：每个供应商的所有药品录入完毕后，输出一行贷方分录，科目 2202 应付账款，
     金额为该供应商所有药品进价金额之和，填写供应商辅助编码与名称；
4. 严格按照 10 项规则填充《表格迁账参考模板.xlsx》的【凭证】Sheet；
5. 自动进行借贷平衡核验及未匹配药品审计。
"""

import os
import sys
import re
import json
import calendar
import argparse
from datetime import datetime, date
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def normalize_text(text):
    """文本标准化：去除所有空格，统一半角全角符号与括号"""
    if text is None:
        return ''
    s = str(text).strip()
    s = re.sub(r'\s+', '', s)
    s = s.replace('（', '(').replace('）', ')')
    s = s.replace('【', '[').replace('】', ']')
    s = s.replace('×', '*').replace('x', '*').replace('X', '*')
    return s


def clean_drug_name(text):
    """去除商品名括号，提取通用名核心（例如 阿立哌唑口崩片（博思清） -> 阿立哌唑口崩片）"""
    norm = normalize_text(text)
    return re.sub(r'\(.*?\)|\[.*?\]', '', norm)


def detect_month_end_date(inbound_file=None, max_inbound_date=None):
    """
    自动检测目标年月并返回月末最后一天的 date 对象
    """
    if max_inbound_date:
        m = re.search(r'(\d{4})[/-](\d{1,2})[/-](\d{1,2})', str(max_inbound_date))
        if m:
            year, month = int(m.group(1)), int(m.group(2))
            last_day = calendar.monthrange(year, month)[1]
            return date(year, month, last_day)

    if inbound_file:
        base_name = os.path.basename(inbound_file)
        m = re.search(r'(\d{4})[^\d]*?(\d{1,2})月?', base_name)
        if m:
            year, month = int(m.group(1)), int(m.group(2))
            if 1 <= month <= 12:
                last_day = calendar.monthrange(year, month)[1]
                return date(year, month, last_day)

    today = date.today()
    last_day = calendar.monthrange(today.year, today.month)[1]
    return date(today.year, today.month, last_day)


def load_factory_abbreviations(config_file='factory_mapping.json'):
    """
    读取并补充厂家全称与简写映射字典
    """
    fmap = {}
    if os.path.exists(config_file):
        try:
            with open(config_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                fmap.update(data.get('factory_abbreviations', {}))
        except Exception as e:
            print(f"[!] 读取 {config_file} 失败: {e}")

    # 补充外账存货辅助名称中常见厂家简写
    extra_abbrs = {
        '沈阳华泰药物研究有限': '华泰',
        '沈阳华泰药物研究有限公司': '华泰',
        '乌兰浩特中蒙制药有限': '乌兰浩特中蒙制药',
        '乌兰浩特中蒙制药有限公司': '乌兰浩特中蒙制药',
        '广东东阳光药业有限公': '广东东阳光',
        '广东东阳光药业有限公司': '广东东阳光',
        '石药集团欧意药业有限公司': '欧意',
        '河北龙海药业有限公司': '河北龙海',
        '苏州第三制药厂有限责': '苏州第三制药厂',
        '苏州第三制药厂有限责任公司': '苏州第三制药厂',
        '合肥立方制药股分有限': '合肥立方',
        '合肥立方制药股份有限公司': '合肥立方',
        '扬子江药业集团上海海': '扬子江',
        '湖南省湘中制药有限公司': '湘中',
        '丽珠集团丽珠制药厂': '丽珠',
        '浙江佐力药业股份有限': '佐力',
        '浙江佐力药业股份有限公司': '佐力',
        'Organon Pharma(UK) Limited': '瑞美隆',
        '湖南洞庭药业股份有限公司': '洞庭',
        '河北国泰医药有限公司': '国泰',
        '国药乐仁堂石家庄医药有限公司': '乐仁堂',
        '华润河北益生医药有限公司': '益生',
        '石药集团河北中诚医药有限公司': '中诚'
    }
    for k, v in extra_abbrs.items():
        if k not in fmap:
            fmap[k] = v

    return fmap


def load_external_aux_data(template_file):
    """
    从《表格迁账参考模板.xlsx》中读取【辅助信息】Sheet，解析供应商与存货字典
    返回: (suppliers, inventory_list)
    """
    if not os.path.exists(template_file):
        raise FileNotFoundError(f"未找到外账模板文件: {template_file}")

    wb = openpyxl.load_workbook(template_file, data_only=True)
    if '辅助信息' not in wb.sheetnames:
        raise ValueError(f"模板文件 {template_file} 缺少【辅助信息】Sheet")

    sheet = wb['辅助信息']
    suppliers = []
    inventory_list = []

    # 表头位于第 2 行，数据行从第 4 行开始（第 3 行为空行）
    for r in list(sheet.iter_rows(values_only=True))[3:]:
        if not r or not any(r):
            continue
        aux_type = str(r[0] or '').strip()
        code = str(r[1] or '').strip()
        name = str(r[2] or '').strip()
        spec = str(r[4] or '').strip() if len(r) > 4 and r[4] is not None else ''
        unit = str(r[5] or '').strip() if len(r) > 5 and r[5] is not None else ''

        if not code or not name:
            continue

        item = {
            'type': aux_type,
            'code': code,
            'name': name,
            'norm_name': normalize_text(name),
            'clean_name': clean_drug_name(name),
            'spec': normalize_text(spec),
            'unit': normalize_text(unit)
        }

        if aux_type == '供应商':
            suppliers.append(item)
        elif aux_type == '存货':
            inventory_list.append(item)

    print(f"[+] 成功加载外账基准档案: 供应商 {len(suppliers)} 家，存货项目 {len(inventory_list)} 种")
    return suppliers, inventory_list


def match_supplier(sup_in, suppliers):
    """
    匹配供应商辅助编码与标准名称
    """
    if not sup_in:
        return None

    norm_in = normalize_text(sup_in)

    # 1. 严格名称匹配
    for s in suppliers:
        if s['norm_name'] == norm_in:
            return s

    # 2. 互相包含匹配
    for s in suppliers:
        if s['norm_name'] in norm_in or norm_in in s['norm_name']:
            return s

    # 3. 去除通用公司后缀后比对
    clean_in = re.sub(r'(股份有限公司|医药有限公司|药业有限公司|有限公司|责任公司|有限责任公司)', '', norm_in)
    for s in suppliers:
        clean_s = re.sub(r'(股份有限公司|医药有限公司|药业有限公司|有限公司|责任公司|有限责任公司)', '', s['norm_name'])
        if clean_in and clean_s and (clean_in in clean_s or clean_s in clean_in):
            return s

    # 4. 关键关键词匹配
    keywords = ['乐仁堂', '益生', '洞庭', '中诚', '国泰', '蕴德', '国瑞堂', '盛方', '神农']
    for kw in keywords:
        if kw in norm_in:
            for s in suppliers:
                if kw in s['norm_name']:
                    return s

    return None


def match_drug(drug_item, inventory_list, factory_abbrs):
    """
    匹配存货辅助编码与标准名称
    """
    name = drug_item['name']
    factory = drug_item.get('factory') or ''
    spec = drug_item.get('spec') or ''

    norm_n = normalize_text(name)
    clean_n = clean_drug_name(name)
    norm_f = normalize_text(factory)
    norm_sp = normalize_text(spec)

    # 查找厂家简写
    abbr = None
    for full_f, short_f in factory_abbrs.items():
        n_full = normalize_text(full_f)
        if n_full == norm_f or n_full in norm_f or norm_f in n_full:
            abbr = short_f
            break

    # 1. 规范名称完全相同
    for item in inventory_list:
        if item['norm_name'] == norm_n:
            return item

    # 2. 带有厂家简写优先匹配（如 氯硝西泮（恩华）、阿立哌唑片（华海5mg））
    if abbr:
        norm_abbr = normalize_text(abbr)
        # 2.1 药名 + 厂家简写 + 规格匹配
        for item in inventory_list:
            if (clean_n in item['clean_name'] or item['clean_name'] in clean_n) and norm_abbr in item['norm_name']:
                if norm_sp and item['spec'] and (norm_sp in item['spec'] or item['spec'] in norm_sp):
                    return item
        # 2.2 药名 + 厂家简写匹配
        for item in inventory_list:
            if (clean_n in item['clean_name'] or item['clean_name'] in clean_n) and norm_abbr in item['norm_name']:
                return item

    # 3. 通用名相同且规格匹配
    for item in inventory_list:
        if item['clean_name'] == clean_n:
            if norm_sp and item['spec'] and (norm_sp in item['spec'] or item['spec'] in norm_sp):
                return item

    # 4. 通用名完全相同
    for item in inventory_list:
        if item['clean_name'] == clean_n:
            return item

    # 5. 模糊子串匹配（长度>=3）
    if len(clean_n) >= 3:
        for item in inventory_list:
            if clean_n in item['clean_name'] or item['clean_name'] in clean_n:
                return item

    return None


def parse_inbound_file(inbound_file):
    """
    读取西药入库单，返回入库明细记录列表与检测到的最大入库日期
    """
    if not os.path.exists(inbound_file):
        raise FileNotFoundError(f"未找到入库单文件: {inbound_file}")

    ext = os.path.splitext(inbound_file)[1].lower()
    inbound_items = []
    max_date_str = None

    if ext == '.xlsx':
        wb = openpyxl.load_workbook(inbound_file, data_only=True)
        # 寻找包含明细的有效 sheet
        target_sheet = None
        header_row = 4
        for s in wb.worksheets:
            for r in range(1, min(10, s.max_row + 1)):
                vals = [str(s.cell(r, c).value or '').strip() for c in range(1, min(25, s.max_column + 1))]
                if '药品名称' in vals and ('数量' in vals or '进价' in vals):
                    target_sheet = s
                    header_row = r
                    break
            if target_sheet:
                break
        if not target_sheet:
            target_sheet = wb.worksheets[0]

        headers = [str(target_sheet.cell(header_row, c).value or '').strip() for c in range(1, target_sheet.max_column + 1)]
        
        def find_c(kw):
            for idx, val in enumerate(headers, start=1):
                if kw in val:
                    return idx
            return -1

        col_name = find_c('药品名称')
        col_factory = find_c('厂家')
        if col_factory < 0: col_factory = find_c('制药厂')
        if col_factory < 0: col_factory = find_c('生产企业')
        col_spec = find_c('规格')
        col_unit = find_c('单位')
        col_qty = find_c('数量')
        col_price = find_c('进价')
        col_amt = find_c('进价金额')
        col_supplier = find_c('供应商')
        col_date = find_c('入库日期')

        for r in range(header_row + 1, target_sheet.max_row + 1):
            raw_name = str(target_sheet.cell(r, col_name).value or '').strip() if col_name > 0 else ''
            if not raw_name or '合计' in raw_name or '制表' in raw_name or raw_name.startswith('报表'):
                continue

            supplier = str(target_sheet.cell(r, col_supplier).value or '').strip() if col_supplier > 0 else ''
            factory = str(target_sheet.cell(r, col_factory).value or '').strip() if col_factory > 0 else ''
            spec = str(target_sheet.cell(r, col_spec).value or '').strip() if col_spec > 0 else ''
            unit = str(target_sheet.cell(r, col_unit).value or '').strip() if col_unit > 0 else ''
            date_val = str(target_sheet.cell(r, col_date).value or '').strip() if col_date > 0 else ''

            if date_val and (max_date_str is None or date_val > max_date_str):
                max_date_str = date_val

            try:
                raw_qty = float(target_sheet.cell(r, col_qty).value or 0.0) if col_qty > 0 else 0.0
                qty = int(raw_qty) if raw_qty.is_integer() else raw_qty
            except (ValueError, TypeError):
                qty = 0.0

            try:
                price = float(target_sheet.cell(r, col_price).value or 0.0) if col_price > 0 else 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                amt = float(target_sheet.cell(r, col_amt).value or 0.0) if col_amt > 0 else 0.0
            except (ValueError, TypeError):
                amt = round(qty * price, 2)

            inbound_items.append({
                'row_index': r,
                'name': raw_name,
                'factory': factory or supplier,
                'spec': spec,
                'unit': unit,
                'qty': qty,
                'price': price,
                'amount': round(amt, 2),
                'supplier': supplier,
                'date': date_val
            })

    elif ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(inbound_file)
        target_sheet = wb.sheet_by_index(0)
        header_row = 3
        for r in range(min(10, target_sheet.nrows)):
            vals = [str(target_sheet.cell_value(r, c)).strip() for c in range(target_sheet.ncols)]
            if '药品名称' in vals and ('数量' in vals or '进价' in vals):
                header_row = r
                break

        headers = [str(target_sheet.cell_value(header_row, c)).strip() for c in range(target_sheet.ncols)]
        def find_xls_c(kw):
            for idx, val in enumerate(headers):
                if kw in val:
                    return idx
            return -1

        col_name = find_xls_c('药品名称')
        col_factory = find_xls_c('厂家')
        if col_factory < 0: col_factory = find_xls_c('制药厂')
        if col_factory < 0: col_factory = find_xls_c('生产企业')
        col_spec = find_xls_c('规格')
        col_unit = find_xls_c('单位')
        col_qty = find_xls_c('数量')
        col_price = find_xls_c('进价')
        col_amt = find_xls_c('进价金额')
        col_supplier = find_xls_c('供应商')
        col_date = find_xls_c('入库日期')

        for r in range(header_row + 1, target_sheet.nrows):
            raw_name = str(target_sheet.cell_value(r, col_name)).strip() if col_name >= 0 else ''
            if not raw_name or '合计' in raw_name or '制表' in raw_name:
                continue

            supplier = str(target_sheet.cell_value(r, col_supplier)).strip() if col_supplier >= 0 else ''
            factory = str(target_sheet.cell_value(r, col_factory)).strip() if col_factory >= 0 else ''
            spec = str(target_sheet.cell_value(r, col_spec)).strip() if col_spec >= 0 else ''
            unit = str(target_sheet.cell_value(r, col_unit)).strip() if col_unit >= 0 else ''
            date_val = str(target_sheet.cell_value(r, col_date)).strip() if col_date >= 0 else ''

            if date_val and (max_date_str is None or date_val > max_date_str):
                max_date_str = date_val

            try:
                raw_qty = float(target_sheet.cell_value(r, col_qty) or 0.0) if col_qty >= 0 else 0.0
                qty = int(raw_qty) if raw_qty.is_integer() else raw_qty
            except (ValueError, TypeError):
                qty = 0.0

            try:
                price = float(target_sheet.cell_value(r, col_price) or 0.0) if col_price >= 0 else 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                amt = float(target_sheet.cell_value(r, col_amt) or 0.0) if col_amt >= 0 else 0.0
            except (ValueError, TypeError):
                amt = round(qty * price, 2)

            inbound_items.append({
                'row_index': r + 1,
                'name': raw_name,
                'factory': factory or supplier,
                'spec': spec,
                'unit': unit,
                'qty': qty,
                'price': price,
                'amount': round(amt, 2),
                'supplier': supplier,
                'date': date_val
            })
    else:
        raise ValueError(f"不支持的入库单格式: {ext}")

    print(f"[+] 成功读取入库单明细: {len(inbound_items)} 行，最新单据日期: {max_date_str}")
    return inbound_items, max_date_str


def build_voucher_entries(inbound_items, suppliers_dict, inventory_list, factory_abbrs, voucher_date=None, voucher_no=1):
    """
    根据用户指定的 10 项规则生成完整的凭证行数据：
    1. 记账日期：当前月末最后一天
    2. 凭证类型：固定“记”
    3. 凭证号：固定 1
    4. 摘要：供应商 + 到货
    5. 科目编码：1405 (借方药品) / 2202 (贷方应付)
    6. 科目名称：库存商品 (借方) / 应付账款 (贷方)
    7. 本币借方金额：入库进价金额
    8. 本币贷方金额：该供应商所有药品录入完毕后的汇总
    9. 数量：入库数量（借方填写，贷方汇总留空）
    10. 辅助编码和辅助名称：从辅助信息取值
    """
    # 按供应商分组保持单据原始出现顺序
    grouped_by_supplier = OrderedDict()
    for item in inbound_items:
        sup = item['supplier'] or '未知供应商'
        if sup not in grouped_by_supplier:
            grouped_by_supplier[sup] = []
        grouped_by_supplier[sup].append(item)

    voucher_rows = []
    unmatched_drugs = []
    unmatched_suppliers = []

    total_debit = 0.0
    total_credit = 0.0

    date_val = voucher_date or date.today()

    for sup_name, items in grouped_by_supplier.items():
        # 1. 匹配该供应商信息
        matched_sup = match_supplier(sup_name, suppliers_dict)
        if matched_sup:
            sup_code = matched_sup['code']
            sup_std_name = matched_sup['name']
        else:
            sup_code = ''
            sup_std_name = sup_name
            unmatched_suppliers.append(sup_name)

        summary_text = f"{sup_std_name}到货"
        supplier_subtotal = 0.0

        # 2. 依次生成该供应商名下药品的【借方分录】
        for d in items:
            matched_drug = match_drug(d, inventory_list, factory_abbrs)
            if matched_drug:
                drug_code = matched_drug['code']
                drug_std_name = matched_drug['name']
            else:
                drug_code = ''
                drug_std_name = d['name']
                unmatched_drugs.append(d)

            amount = round(d['amount'], 2)
            supplier_subtotal += amount
            total_debit += amount

            # 借方分录行 (15 列)
            voucher_rows.append({
                'date': date_val,
                'voucher_type': '记',
                'voucher_no': voucher_no,
                'summary': summary_text,
                'subject_code': 1405,
                'subject_name': '库存商品',
                'debit_amount': amount,
                'credit_amount': None,
                'foreign_debit': None,
                'foreign_credit': None,
                'qty': d['qty'],
                'aux_code': drug_code,
                'aux_name': drug_std_name,
                'attachment_num': None,
                'creator': None,
                'raw_item': d
            })

        supplier_subtotal = round(supplier_subtotal, 2)
        total_credit += supplier_subtotal

        # 3. 该供应商药品录入完成后，紧跟一行【贷方汇总分录】
        voucher_rows.append({
            'date': date_val,
            'voucher_type': '记',
            'voucher_no': voucher_no,
            'summary': summary_text,
            'subject_code': 2202,
            'subject_name': '应付账款',
            'debit_amount': None,
            'credit_amount': supplier_subtotal,
            'foreign_debit': None,
            'foreign_credit': None,
            'qty': None,
            'aux_code': sup_code,
            'aux_name': sup_std_name,
            'attachment_num': None,
            'creator': None,
            'raw_item': None
        })

    total_debit = round(total_debit, 2)
    total_credit = round(total_credit, 2)
    diff = round(total_debit - total_credit, 2)

    return voucher_rows, {
        'total_debit': total_debit,
        'total_credit': total_credit,
        'diff': diff,
        'unmatched_drugs': unmatched_drugs,
        'unmatched_suppliers': unmatched_suppliers,
        'supplier_count': len(grouped_by_supplier),
        'item_count': len(inbound_items)
    }


def write_external_voucher_to_excel(template_file, output_file, voucher_rows):
    """
    将生成的凭证分录行写入《表格迁账参考模板.xlsx》的【凭证】Sheet
    保留第 1~3 行的原有表头与样式，从第 4 行开始写入，并设置专业的会计格式。
    """
    wb = openpyxl.load_workbook(template_file)
    if '凭证' not in wb.sheetnames:
        raise ValueError("模板文件中未找到【凭证】工作表！")

    sheet = wb['凭证']

    # 清空第 4 行及以后的旧示例数据
    max_r = sheet.max_row
    if max_r >= 4:
        for r in range(4, max_r + 1):
            for c in range(1, 16):
                cell = sheet.cell(row=r, column=c)
                cell.value = None
                cell.border = None
                cell.fill = PatternFill(fill_type=None)

    # 样式定义
    font_main = Font(name='微软雅黑', size=10)
    font_bold = Font(name='微软雅黑', size=10, bold=True)
    align_left = Alignment(horizontal='left', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    thin_border_side = Side(border_style='thin', color='D9D9D9')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    credit_fill = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid') # 浅绿提示贷方汇总
    warn_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')   # 浅黄提示未匹配

    current_row = 4
    for entry in voucher_rows:
        is_credit_row = (entry['credit_amount'] is not None)
        is_unmatched = (not entry['aux_code'])

        row_values = [
            entry['date'],
            entry['voucher_type'],
            entry['voucher_no'],
            entry['summary'],
            entry['subject_code'],
            entry['subject_name'],
            entry['debit_amount'],
            entry['credit_amount'],
            entry['foreign_debit'],
            entry['foreign_credit'],
            entry['qty'],
            entry['aux_code'],
            entry['aux_name'],
            entry['attachment_num'],
            entry['creator']
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = sheet.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_bold if is_credit_row else font_main
            cell.border = cell_border

            # 背景高亮
            if is_unmatched and col_idx in [12, 13]:
                cell.fill = warn_fill
            elif is_credit_row:
                cell.fill = credit_fill

            # 对齐方式与格式化
            if col_idx in [1, 2, 3]:
                cell.alignment = align_center
                if col_idx == 1 and isinstance(val, (date, datetime)):
                    cell.number_format = 'yyyy-mm-dd'
            elif col_idx in [7, 8]:
                cell.alignment = align_right
                if val is not None:
                    cell.number_format = '#,##0.00'
            elif col_idx == 11:
                cell.alignment = align_right
                if val is not None:
                    cell.number_format = '#,##0.##'
            elif col_idx == 12:
                cell.alignment = align_center
                cell.number_format = '@' # 文本格式保证 5 位纯数字 00028 不丢失前导 0
            else:
                cell.alignment = align_left

        current_row += 1

    # 清理多余空行，确保导出的表格行数与分录行数完全吻合
    if sheet.max_row >= current_row:
        sheet.delete_rows(current_row, sheet.max_row - current_row + 1)

    wb.save(output_file)
    print(f"[✓] 成功生成外账凭证文件: {output_file}")


def process_external_inbound(inbound_file='西药入库单.xlsx',
                             template_file='表格迁账参考模板.xlsx',
                             output_file=None,
                             config_file='factory_mapping.json',
                             custom_date=None,
                             voucher_no=1,
                             as_json=False):
    """
    执行完整的西药外账入库凭证生成流程
    """
    if not as_json:
        print("=" * 65)
        print("🏥 石家庄心理医院 · 外账入库凭证自动化生成引擎")
        print("=" * 65)
        print(f"[*] 入库单文件: {inbound_file}")
        print(f"[*] 外账参考模板: {template_file}")

    # 1. 加载档案与配置
    factory_abbrs = load_factory_abbreviations(config_file)
    suppliers_dict, inventory_list = load_external_aux_data(template_file)

    # 2. 读取入库明细
    inbound_items, max_date = parse_inbound_file(inbound_file)
    if custom_date:
        parts = str(custom_date).strip().split('-')
        if len(parts) == 3:
            voucher_date = date(int(parts[0]), int(parts[1]), int(parts[2]))
        else:
            voucher_date = detect_month_end_date(inbound_file=inbound_file, max_inbound_date=max_date)
    else:
        voucher_date = detect_month_end_date(inbound_file=inbound_file, max_inbound_date=max_date)

    if not as_json:
        print(f"[*] 确定记账日期: {voucher_date} (月末最后一天)")

    # 3. 组织凭证分录
    try:
        v_no = int(voucher_no)
    except (ValueError, TypeError):
        v_no = 1

    voucher_rows, audit_stats = build_voucher_entries(
        inbound_items,
        suppliers_dict,
        inventory_list,
        factory_abbrs,
        voucher_date=voucher_date,
        voucher_no=v_no
    )

    # 4. 打印审计与核验信息
    if not as_json:
        print("-" * 65)
        print("📊 凭证分录与借贷平衡审计报告:")
        print(f"  • 入库单药品记录数: {audit_stats['item_count']} 行")
        print(f"  • 涉及供货商数量  : {audit_stats['supplier_count']} 家")
        print(f"  • 生成凭证总分录数: {len(voucher_rows)} 行 (含借方明细 + 贷方汇总)")
        print(f"  • 借方总金额 (1405): ¥{audit_stats['total_debit']:,.2f}")
        print(f"  • 贷方总金额 (2202): ¥{audit_stats['total_credit']:,.2f}")
        print(f"  • 借贷平衡差额    : ¥{audit_stats['diff']:,.2f} ({'✅ 完全平衡' if audit_stats['diff'] == 0 else '❌ 借贷不平'})")

        unmatched_drugs = audit_stats['unmatched_drugs']
        if unmatched_drugs:
            print(f"\n⚠️  存在 {len(unmatched_drugs)} 笔药品未匹配到外账 5 位存货辅助编码 (已在 Excel 中高亮标黄):")
            for idx, d in enumerate(unmatched_drugs, 1):
                print(f"   {idx}. [行{d['row_index']}] {d['name']} | 规格: {d['spec']} | 厂家: {d['factory']} | 数量: {d['qty']} | 金额: ¥{d['amount']}")
        else:
            print("\n✅ 所有药品存货辅助编码全部匹配成功！")

    # 5. 写入目标 Excel
    if not output_file:
        output_file = '表格迁账参考模板_西药入库_已生成.xlsx'

    write_external_voucher_to_excel(template_file, output_file, voucher_rows)
    if not as_json:
        print("=" * 65)

    if as_json:
        total_items = len(inbound_items)
        unmatched_cnt = len(audit_stats['unmatched_drugs'])
        matched_cnt = total_items - unmatched_cnt
        match_rate = round((matched_cnt / total_items * 100.0) if total_items > 0 else 100.0, 2)

        # 构造可序列化的前瞻数据行
        serializable_rows = []
        for idx, r in enumerate(voucher_rows, 1):
            serializable_rows.append({
                'row_no': idx,
                'date': str(r['date']),
                'voucher_type': r['voucher_type'],
                'voucher_no': r['voucher_no'],
                'summary': r['summary'],
                'subject_code': r['subject_code'],
                'subject_name': r['subject_name'],
                'debit_amount': r['debit_amount'],
                'credit_amount': r['credit_amount'],
                'qty': r['qty'],
                'aux_code': r['aux_code'],
                'aux_name': r['aux_name'],
                'is_credit': (r['credit_amount'] is not None),
                'is_unmatched': (not r['aux_code'])
            })

        json_result = {
            'success': True,
            'output_file': os.path.abspath(output_file),
            'voucher_date': str(voucher_date),
            'voucher_no': v_no,
            'total_items': total_items,
            'supplier_count': audit_stats['supplier_count'],
            'total_entries': len(voucher_rows),
            'matched_count': matched_cnt,
            'unmatched_count': unmatched_cnt,
            'match_rate': match_rate,
            'total_debit': audit_stats['total_debit'],
            'total_credit': audit_stats['total_credit'],
            'diff': audit_stats['diff'],
            'is_balanced': (audit_stats['diff'] == 0.0),
            'unmatched_drugs': [
                {
                    'row_index': d['row_index'],
                    'name': d['name'],
                    'factory': d['factory'],
                    'spec': d['spec'],
                    'qty': d['qty'],
                    'amount': d['amount'],
                    'supplier': d['supplier']
                } for d in audit_stats['unmatched_drugs']
            ],
            'voucher_rows': serializable_rows
        }
        print(json.dumps(json_result, ensure_ascii=False))

    return output_file, audit_stats, voucher_rows


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="根据西药入库单自动生成外账迁账凭证")
    parser.add_argument('-i', '--inbound', default='西药入库单.xlsx', help="药品入库单文件路径")
    parser.add_argument('-t', '--template', default='表格迁账参考模板.xlsx', help="外账参考模板文件路径")
    parser.add_argument('-o', '--output', default='表格迁账参考模板_西药入库_已生成.xlsx', help="输出 Excel 文件路径")
    parser.add_argument('-c', '--config', default='factory_mapping.json', help="厂家映射配置文件路径")
    parser.add_argument('-d', '--date', default=None, help="自定义记账日期 YYYY-MM-DD")
    parser.add_argument('-v', '--voucher-no', default=1, help="凭证号 (默认 1)")
    parser.add_argument('--json', action='store_true', help="输出结构化 JSON 结果供桌面客户端解析")

    args = parser.parse_args()
    process_external_inbound(
        inbound_file=args.inbound,
        template_file=args.template,
        output_file=args.output,
        config_file=args.config,
        custom_date=args.date,
        voucher_no=args.voucher_no,
        as_json=args.json
    )
