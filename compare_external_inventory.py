#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
外账账实库存智能核对脚本 (compare_external_inventory.py)
功能：
1. 读取外账导入模板（表格迁账参考模板.xlsx）中的【辅助余额表】1405 存货科目；
2. 读取库管系统导出的【西药 / 中药 / 耗材 库存汇总报表】（.xls 或 .xlsx）；
3. 结合厂家简写与规格清洗，精准比对外账账面结存与库管实盘在库数量；
4. 输出四维状态：完全吻合、数量差异、仅外账有、仅库管有；
5. 支持输出结构化 JSON 供桌面客户端展示，并可导出《外账账实库存核对分析报告.xlsx》。
"""

import os
import sys
import re
import json
import argparse
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

try:
    import xlrd
except ImportError:
    xlrd = None


def normalize_text(text):
    """文本标准化：去除所有空格，统一半角全角符号与括号"""
    if text is None:
        return ''
    s = str(text).strip()
    s = re.sub(r'\s+', '', s)
    s = s.replace('（', '(').replace('）', ')')
    s = s.replace('【', '[').replace('】', ']')
    s = s.replace('×', '*').replace('x', '*').replace('X', '*')
    return s


def clean_drug_name(text):
    """去除商品名括号，提取通用名核心"""
    norm = normalize_text(text)
    return re.sub(r'\(.*?\)|\[.*?\]', '', norm)


def load_factory_abbreviations(config_file='factory_mapping.json'):
    """读取厂家简写映射"""
    fmap = {}
    if os.path.exists(config_file):
        try:
            with open(config_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                fmap.update(data.get('factory_abbreviations', {}))
        except Exception:
            pass

    extra = {
        '沈阳华泰药物研究有限': '华泰',
        '沈阳华泰药物研究有限公司': '华泰',
        '乌兰浩特中蒙制药有限': '乌兰浩特中蒙制药',
        '乌兰浩特中蒙制药有限公司': '乌兰浩特中蒙制药',
        '广东东阳光药业有限公': '广东东阳光',
        '广东东阳光药业有限公司': '广东东阳光',
        '石药集团欧意药业有限公司': '欧意',
        '河北龙海药业有限公司': '河北龙海',
        '苏州第三制药厂有限责': '苏州第三制药厂',
        '苏州第三制药厂有限责任公司': '苏州第三制药厂',
        '合肥立方制药股分有限': '合肥立方',
        '扬子江药业集团上海海': '扬子江',
        '湖南省湘中制药有限公司': '湘中',
        '丽珠集团丽珠制药厂': '丽珠',
        '浙江佐力药业股份有限': '佐力',
        'Organon Pharma(UK) Limited': '瑞美隆',
        '湖南洞庭药业股份有限公司': '洞庭',
        '河北国泰医药有限公司': '国泰',
        '国药乐仁堂石家庄医药有限公司': '乐仁堂',
        '华润河北益生医药有限公司': '益生',
        '石药集团河北中诚医药有限公司': '中诚'
    }
    for k, v in extra.items():
        if k not in fmap:
            fmap[k] = v
    return fmap


def load_external_balance_items(template_file):
    """
    从《表格迁账参考模板.xlsx》的【辅助余额表】中读取 1405 存货科目的期初/结存数量
    """
    if not os.path.exists(template_file):
        raise FileNotFoundError(f"未找到外账模板文件: {template_file}")

    wb = openpyxl.load_workbook(template_file, data_only=True)
    if '辅助余额表' not in wb.sheetnames:
        raise ValueError(f"模板文件 {template_file} 缺少【辅助余额表】工作表")

    sheet = wb['辅助余额表']
    items = []
    # 数据行从第 4 行开始
    for r in list(sheet.iter_rows(values_only=True))[3:]:
        if not r or not any(r):
            continue
        subj_code = str(r[0] or '').strip()
        aux_type = str(r[2] or '').strip()
        code = str(r[3] or '').strip()
        name = str(r[4] or '').strip()
        unit = str(r[5] or '').strip() if len(r) > 5 and r[5] is not None else ''

        if subj_code == '1405' and aux_type == '存货' and code and name:
            def parse_num(val):
                try:
                    return float(val) if val is not None and str(val).strip() != '' else 0.0
                except (ValueError, TypeError):
                    return 0.0

            qty = parse_num(r[8]) if len(r) > 8 else 0.0
            amt = parse_num(r[10]) if len(r) > 10 else 0.0

            items.append({
                'code': code,
                'name': name,
                'norm_name': normalize_text(name),
                'clean_name': clean_drug_name(name),
                'unit': unit,
                'ext_qty': qty,
                'ext_amt': amt
            })

    return items


def load_warehouse_items(warehouse_file):
    """
    读取库管系统在库报表（.xlsx 或 .xls）
    """
    if not os.path.exists(warehouse_file):
        raise FileNotFoundError(f"未找到库管文件: {warehouse_file}")

    ext = os.path.splitext(warehouse_file)[1].lower()
    wh_items = []

    if ext == '.xls':
        if not xlrd:
            raise ImportError("读取 .xls 文件需要 xlrd 库")
        wb = xlrd.open_workbook(warehouse_file)
        sheet = wb.sheet_by_index(0)

        header_row = 0
        for r in range(min(10, sheet.nrows)):
            vals = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
            if '药品名称' in vals and '数量' in vals:
                header_row = r
                break

        headers = [str(sheet.cell_value(header_row, c)).strip() for c in range(sheet.ncols)]
        def find_c(kw):
            for idx, h in enumerate(headers):
                if kw in h: return idx
            return -1

        c_name = find_c('药品名称')
        c_spec = find_c('规格')
        c_factory = find_c('制药厂')
        if c_factory < 0: c_factory = find_c('生产企业')
        if c_factory < 0: c_factory = find_c('厂家')
        c_unit = find_c('单位')
        c_qty = find_c('数量')
        c_price = find_c('单价')
        c_amt = find_c('金额')

        for r in range(header_row + 1, sheet.nrows):
            name = str(sheet.cell_value(r, c_name)).strip() if c_name >= 0 else ''
            if not name or '合计' in name or '制表' in name:
                continue

            spec = str(sheet.cell_value(r, c_spec)).strip() if c_spec >= 0 else ''
            factory = str(sheet.cell_value(r, c_factory)).strip() if c_factory >= 0 else ''
            unit = str(sheet.cell_value(r, c_unit)).strip() if c_unit >= 0 else ''

            try:
                qty = float(sheet.cell_value(r, c_qty) or 0.0) if c_qty >= 0 else 0.0
            except (ValueError, TypeError):
                qty = 0.0

            try:
                price = float(sheet.cell_value(r, c_price) or 0.0) if c_price >= 0 else 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                amt = float(sheet.cell_value(r, c_amt) or 0.0) if c_amt >= 0 else 0.0
            except (ValueError, TypeError):
                amt = round(qty * price, 2)

            wh_items.append({
                'row_no': r + 1,
                'name': name,
                'spec': spec,
                'factory': factory,
                'unit': unit,
                'wh_qty': qty,
                'price': price,
                'wh_amt': amt,
                'norm_name': normalize_text(name),
                'clean_name': clean_drug_name(name)
            })

    elif ext == '.xlsx':
        wb = openpyxl.load_workbook(warehouse_file, data_only=True)
        sheet = wb.worksheets[0]

        header_row = 1
        for r in range(1, min(10, sheet.max_row + 1)):
            vals = [str(sheet.cell(r, c).value or '').strip() for c in range(1, min(20, sheet.max_column + 1))]
            if '药品名称' in vals and '数量' in vals:
                header_row = r
                break

        headers = [str(sheet.cell(header_row, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
        def find_c(kw):
            for idx, h in enumerate(headers, 1):
                if kw in h: return idx
            return -1

        c_name = find_c('药品名称')
        c_spec = find_c('规格')
        c_factory = find_c('制药厂')
        if c_factory < 0: c_factory = find_c('生产企业')
        if c_factory < 0: c_factory = find_c('厂家')
        c_unit = find_c('单位')
        c_qty = find_c('数量')
        c_price = find_c('单价')
        c_amt = find_c('金额')

        for r in range(header_row + 1, sheet.max_row + 1):
            name = str(sheet.cell(r, c_name).value or '').strip() if c_name > 0 else ''
            if not name or '合计' in name or '制表' in name:
                continue

            spec = str(sheet.cell(r, c_spec).value or '').strip() if c_spec > 0 else ''
            factory = str(sheet.cell(r, c_factory).value or '').strip() if c_factory > 0 else ''
            unit = str(sheet.cell(r, c_unit).value or '').strip() if c_unit > 0 else ''

            try:
                qty = float(sheet.cell(r, c_qty).value or 0.0) if c_qty > 0 else 0.0
            except (ValueError, TypeError):
                qty = 0.0

            try:
                price = float(sheet.cell(r, c_price).value or 0.0) if c_price > 0 else 0.0
            except (ValueError, TypeError):
                price = 0.0

            try:
                amt = float(sheet.cell(r, c_amt).value or 0.0) if c_amt > 0 else 0.0
            except (ValueError, TypeError):
                amt = round(qty * price, 2)

            wh_items.append({
                'row_no': r,
                'name': name,
                'spec': spec,
                'factory': factory,
                'unit': unit,
                'wh_qty': qty,
                'price': price,
                'wh_amt': amt,
                'norm_name': normalize_text(name),
                'clean_name': clean_drug_name(name)
            })

    return wh_items


def compare_external_inventory_data(template_file, warehouse_file, config_file='factory_mapping.json'):
    """
    进行外账账面结存与库管实盘比对
    """
    factory_abbrs = load_factory_abbreviations(config_file)
    ext_items = load_external_balance_items(template_file)
    wh_items = load_warehouse_items(warehouse_file)

    # 建立外账查找索引
    ext_matched_set = set()
    records = []

    equal_count = 0
    diff_count = 0
    ext_only_count = 0
    wh_only_count = 0

    # 1. 遍历库管在库记录，匹配外账账面
    for wh in wh_items:
        norm_n = wh['norm_name']
        clean_n = wh['clean_name']
        norm_f = normalize_text(wh['factory'])

        abbr = None
        for full_f, short_f in factory_abbrs.items():
            n_full = normalize_text(full_f)
            if n_full == norm_f or n_full in norm_f or norm_f in n_full:
                abbr = short_f
                break

        matched_ext = None
        # 1.1 全名精确
        for item in ext_items:
            if item['norm_name'] == norm_n:
                matched_ext = item
                break

        # 1.2 药名 + 厂家简写
        if not matched_ext and abbr:
            norm_abbr = normalize_text(abbr)
            for item in ext_items:
                if (clean_n in item['clean_name'] or item['clean_name'] in clean_n) and norm_abbr in item['norm_name']:
                    matched_ext = item
                    break

        # 1.3 通用名精确
        if not matched_ext:
            for item in ext_items:
                if item['clean_name'] == clean_n:
                    matched_ext = item
                    break

        # 1.4 子串匹配
        if not matched_ext and len(clean_n) >= 3:
            for item in ext_items:
                if clean_n in item['clean_name'] or item['clean_name'] in clean_n:
                    matched_ext = item
                    break

        if matched_ext:
            ext_matched_set.add(matched_ext['code'])
            ext_qty = matched_ext['ext_qty']
            ext_amt = matched_ext['ext_amt']
            diff_qty = round(ext_qty - wh['wh_qty'], 3)
            diff_amt = round(ext_amt - wh['wh_amt'], 2)

            if abs(diff_qty) < 0.001:
                status = 'EQUAL'
                status_label = '完全吻合'
                equal_count += 1
            else:
                status = 'DIFF'
                status_label = '数量差异'
                diff_count += 1

            records.append({
                'ext_code': matched_ext['code'],
                'ext_name': matched_ext['name'],
                'wh_name': wh['name'],
                'wh_spec': wh['spec'],
                'wh_factory': wh['factory'],
                'unit': wh['unit'] or matched_ext['unit'],
                'ext_qty': ext_qty,
                'wh_qty': wh['wh_qty'],
                'diff_qty': diff_qty,
                'ext_amt': ext_amt,
                'wh_amt': wh['wh_amt'],
                'diff_amt': diff_amt,
                'status': status,
                'status_label': status_label
            })
        else:
            wh_only_count += 1
            records.append({
                'ext_code': '',
                'ext_name': '未在账面建档',
                'wh_name': wh['name'],
                'wh_spec': wh['spec'],
                'wh_factory': wh['factory'],
                'unit': wh['unit'],
                'ext_qty': 0.0,
                'wh_qty': wh['wh_qty'],
                'diff_qty': round(-wh['wh_qty'], 3),
                'ext_amt': 0.0,
                'wh_amt': wh['wh_amt'],
                'diff_amt': round(-wh['wh_amt'], 2),
                'status': 'WH_ONLY',
                'status_label': '仅库管有'
            })

    # 2. 检查外账有结存但库管没有的记录
    for ext in ext_items:
        if ext['code'] not in ext_matched_set and (ext['ext_qty'] > 0 or ext['ext_amt'] > 0):
            ext_only_count += 1
            records.append({
                'ext_code': ext['code'],
                'ext_name': ext['name'],
                'wh_name': '-',
                'wh_spec': '-',
                'wh_factory': '-',
                'unit': ext['unit'],
                'ext_qty': ext['ext_qty'],
                'wh_qty': 0.0,
                'diff_qty': ext['ext_qty'],
                'ext_amt': ext['ext_amt'],
                'wh_amt': 0.0,
                'diff_amt': ext['ext_amt'],
                'status': 'EXT_ONLY',
                'status_label': '仅外账有'
            })

    total_items = len(records)
    match_rate = round((equal_count / total_items * 100.0) if total_items > 0 else 0.0, 2)

    return {
        'total_items': total_items,
        'equal_count': equal_count,
        'diff_count': diff_count,
        'ext_only_count': ext_only_count,
        'wh_only_count': wh_only_count,
        'match_rate': match_rate,
        'records': records
    }


def export_audit_excel(audit_data, output_file='外账账实库存核对分析报告.xlsx'):
    """导出核对分析报告 Excel 文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '账实比对底稿'

    headers = [
        '外账辅助编码', '外账存货名称', '库管药品名称', '规格型号',
        '生产厂家', '单位', '外账账面数量', '库管在库数量', '数量差异',
        '外账账面金额', '库管在库金额', '金额差异', '比对状态'
    ]

    font_header = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    font_main = Font(name='微软雅黑', size=10)

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, r in enumerate(audit_data['records'], 2):
        vals = [
            r['ext_code'], r['ext_name'], r['wh_name'], r['wh_spec'],
            r['wh_factory'], r['unit'], r['ext_qty'], r['wh_qty'], r['diff_qty'],
            r['ext_amt'], r['wh_amt'], r['diff_amt'], r['status_label']
        ]
        for col_idx, val in enumerate(vals, 1):
            cell = ws.cell(row_idx, col_idx, val)
            cell.font = font_main
            if col_idx in [1, 6, 13]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in [7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0.00' if col_idx in [10, 11, 12] else '#,##0.##'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # 自动列宽
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_file)
    return output_file


def process_external_inventory_audit(template_file='表格迁账参考模板.xlsx',
                                     warehouse_file='石家庄心理医院新西药房库存汇总报表2026831.xls',
                                     output_file='外账账实库存核对分析报告.xlsx',
                                     config_file='factory_mapping.json',
                                     as_json=False):
    """外账结存数比对主流程"""
    audit_data = compare_external_inventory_data(
        template_file=template_file,
        warehouse_file=warehouse_file,
        config_file=config_file
    )

    if output_file:
        export_audit_excel(audit_data, output_file)
        audit_data['output_file'] = os.path.abspath(output_file)

    if as_json:
        result = {
            'success': True,
            'output_file': audit_data.get('output_file', ''),
            'total_items': audit_data['total_items'],
            'equal_count': audit_data['equal_count'],
            'diff_count': audit_data['diff_count'],
            'ext_only_count': audit_data['ext_only_count'],
            'wh_only_count': audit_data['wh_only_count'],
            'match_rate': audit_data['match_rate'],
            'records': audit_data['records'][:300] # 返回前300条供前端预览
        }
        print(json.dumps(result, ensure_ascii=False))

    return audit_data


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="外账账面结存与库管库存比对")
    parser.add_argument('-t', '--template', default='表格迁账参考模板.xlsx', help="外账参考模板")
    parser.add_argument('-w', '--warehouse', default='石家庄心理医院新西药房库存汇总报表2026831.xls', help="库管在库报表")
    parser.add_argument('-o', '--output', default='外账账实库存核对分析报告.xlsx', help="输出分析报告路径")
    parser.add_argument('-c', '--config', default='factory_mapping.json', help="厂家配置字典")
    parser.add_argument('--json', action='store_true', help="输出结构化 JSON 结果")

    args = parser.parse_args()
    process_external_inventory_audit(
        template_file=args.template,
        warehouse_file=args.warehouse,
        output_file=args.output,
        config_file=args.config,
        as_json=args.json
    )
